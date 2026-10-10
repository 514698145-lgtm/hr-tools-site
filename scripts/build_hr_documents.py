"""Generate editable, print-ready individual Word forms from explicit designs."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from template_designs import WORD, EXCEL
ROOT=Path(__file__).resolve().parents[1]
out=ROOT/'outputs/hr-templates';out.mkdir(exist_ok=True,parents=True)
word_out=out/'word-current';word_out.mkdir(exist_ok=True)
catalog=json.loads((ROOT/'public/data/hr-templates.json').read_text('utf-8'))
refs={r['id']:r for r in json.loads((ROOT/'docs/template-references.json').read_text('utf-8'))}

def shade(cell,fill):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(el)
def table(doc,heads,rows,seeds=None):
    tab=doc.add_table(rows=1,cols=len(heads));tab.autofit=False
    width=6.95/len(heads)
    for c in tab.columns:c.width=Inches(width)
    props=tab._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D4D9DF');borders.append(node)
    props.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        node=OxmlElement('w:'+edge);node.set(qn('w:w'),'100');node.set(qn('w:type'),'dxa');margin.append(node)
    props.append(margin)
    for c,label in zip(tab.rows[0].cells,heads):
        c.text=label;shade(c,'293C50')
        for run in c.paragraphs[0].runs:run.font.color.rgb=RGBColor.from_string('FFFFFF');run.bold=True;run.font.size=Pt(9)
    header=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(header)
    for i in range(rows):
        row=tab.add_row();cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,c in enumerate(row.cells):
            c.text=(seeds[i] if seeds and j==0 and i<len(seeds) else '\u3000')
            c.paragraphs[0].paragraph_format.space_after=Pt(7)
            if i%2==0:shade(c,'F5F7F9')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return tab

manifest=[]
for t in catalog:
    if t['id'] not in WORD:continue
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Noto Sans CJK SC';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(10);doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
    for node in list(doc.styles['Title']._element.iter(qn('w:pBdr'))):node.getparent().remove(node)
    if t['id']=='job-profile':
        doc.styles['Normal'].paragraph_format.space_after=Pt(2)
        doc.styles['Heading 2'].paragraph_format.space_before=Pt(6)
        sec.top_margin=Inches(.5)
    header=sec.header.paragraphs[0];header.text=t['category']+'  /  工作表单';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('747D88')
    doc.add_paragraph(t['title'],'Title')
    p=doc.add_paragraph(t['description']+'。请按实际情况填写，不适用的项目可留空或注明。');p.runs[0].font.size=Pt(9);p.runs[0].font.color.rgb=RGBColor.from_string('596575')
    for block in WORD[t['id']]:
        kind,title,*args=block;doc.add_paragraph(title,'Heading 2')
        if kind=='fields':
            labels=args[0].split('|')
            # Compact 2-column metadata, deliberately borderless with generous writing space.
            for i in range(0,len(labels),2):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(3.55))
                p.add_run(labels[i]+'  ________________')
                if i+1<len(labels):p.add_run('\t'+labels[i+1]+'  ________________')
        elif kind=='writing':
            prompt,lines=args;p=doc.add_paragraph(prompt);p.runs[0].font.size=Pt(9);p.runs[0].font.color.rgb=RGBColor.from_string('667085')
            for _ in range(lines):
                p=doc.add_paragraph(' ');p.paragraph_format.space_after=Pt(8)
                pb=OxmlElement('w:pBdr');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'DEE2E7');pb.append(b);p._p.get_or_add_pPr().append(pb)
        elif kind=='table':table(doc,args[0],args[1],args[2] if len(args)>2 else None)
        elif kind=='checks':doc.add_paragraph('    '.join('□ '+x for x in args[0].split('|')))
        elif kind=='scale':
            doc.add_paragraph('1 很低  ·  2 较低  ·  3 一般  ·  4 较高  ·  5 很高  ·  不适用可留空').runs[0].font.size=Pt(9)
            tab=table(doc,['评价项目','请勾选','事实或补充说明'],len(args[0]),args[0])
            for row in tab.rows[1:]:row.cells[1].text='□1  □2  □3  □4  □5'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('内部工作记录  ·  按授权范围保管  |  ').font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    doc.core_properties.title=t['title'];doc.core_properties.subject=t['description'];doc.core_properties.author='HR 工作站'
    doc.save(word_out/(t['id']+'.docx'))
    fields=[]
    for block in WORD[t['id']]:
        fields.extend(block[2].split('|') if block[0]=='fields' else block[2] if block[0] in ['table','scale'] else [block[1]])
    manifest.append({**t,'fields':list(dict.fromkeys(fields)),'format':'docx','formulaColumns':[],'note':'Word 可编辑表单；填写与签批区域已按本场景独立排版。','sections':[x[1] for x in WORD[t['id']]],'reference':refs[t['id']]['url']})
(out/'word-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),'utf-8')
specs={k:{'layout':v[0],'groups':v[1].split('|'),'sizes':v[2],'tasks':v[3]} for k,v in EXCEL.items()}
(out/'excel-designs.json').write_text(json.dumps(specs,ensure_ascii=False,indent=2),'utf-8')
print('Authored',len(manifest),'individual Word forms')
